"""课件电子化服务 —— 把 PPTX 拆解成「电子书」

设计动机（对应赛题需求 2）：
1. PPTX 动辄几十 MB，整包下载浏览慢 → 拆解成「文本元数据 + 逐页瘦身分图」，
   首屏只需几十 KB 的 JSON，图片按阅读位置懒加载；
2. PPT 不是好的阅读载体（要点破碎、无连贯正文、无法搜索选中）
   → 提取出真正的文本层（标题/要点/表格/备注），支持选中、搜索、
     并可拼装成连贯的「读物模式」。

同时提供 render_pptx() 反向导出，保证老师课堂投影仍能拿到原生 .pptx。
"""
import io
import os
import logging
from typing import Any

from app.utils import paths

logger = logging.getLogger(__name__)

# 图片瘦身参数：长边上限 + JPEG 质量。1400px 足够全屏阅读，
# 一般可把单页压到 60~180KB。
MAX_IMAGE_EDGE = 1400
JPEG_QUALITY = 78
MIN_IMAGE_BYTES = 8 * 1024  # 小于此说明是装饰图形，不值得单独存


def _slides_dir(courseware_id: int) -> tuple[str, str]:
    """课件渲染图落盘目录：固定根 backend/generated/courseware/…

    返回 (相对 GENERATED_ROOT 的入库路径, 绝对目录)，图片是派生产物不是用户上传，
    因此不再写到 uploads/ 里。
    """
    abs_dir = paths.gen_dir("courseware", f"c{courseware_id // 100 or 1}", f"cw{courseware_id}")
    return paths.rel("courseware", f"c{courseware_id // 100 or 1}", f"cw{courseware_id}"), str(abs_dir)


def _compress_image(blob: bytes, out_abs: str, out_rel: str, name: str) -> tuple[str, int]:
    """内嵌图片重压缩落盘，返回 (相对路径, KB)"""
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(blob))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        w, h = img.size
        scale = min(1.0, MAX_IMAGE_EDGE / max(w, h))
        if scale < 1.0:
            img = img.resize((int(w * scale), int(h * scale)), Image.LANCZOS)
        path = os.path.join(out_abs, name)
        img.save(path, "JPEG", quality=JPEG_QUALITY, optimize=True, progressive=True)
        return os.path.join(out_rel, name).replace("\\", "/"), max(1, os.path.getsize(path) // 1024)
    except Exception as e:
        logger.warning(f"image compress failed: {e}")
        return "", 0


def _extract_text_shapes(slide) -> tuple[str, list[dict], list[list[str]]]:
    """提取标题、要点（带层级）、表格"""
    from pptx.util import Emu
    title = ""
    bullets: list[dict] = []
    tables: list[list[str]] = []

    # 占位符顺序：标题优先
    try:
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()
    except Exception:
        pass

    for shape in slide.shapes:
        # 表格
        if getattr(shape, "has_table", False):
            try:
                rows = []
                for row in shape.table.rows:
                    rows.append([(c.text or "").strip() for c in row.cells])
                tables.append(rows)
            except Exception:
                pass
            continue
        # 文本框
        if not getattr(shape, "has_text_frame", False):
            continue
        # 跳过标题占位符（idx==0）——注意 slide.shapes.title 每次访问返回新代理对象，
        # 用 `is` 比较永假，必须按 placeholder idx 判定
        if _is_title_placeholder(shape):
            continue
        tf = shape.text_frame
        for para in tf.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            level = int(para.level or 0)
            if not title and level == 0 and len(text) < 120:
                # 无标题占位符时，取第一行短文本当标题
                title = text
                continue
            bullets.append({"level": level, "text": text})
    return title, bullets[:40], tables[:6]


def _is_title_placeholder(shape) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        return int(shape.placeholder_format.idx) == 0
    except Exception:
        return False


def _draw_slide_image(title: str, bullets: list[dict], out_abs: str, out_rel: str,
                      name: str, page_no: int = 0) -> tuple[str, int]:
    """版式化预览图（委托 deck_design，保证与导出 pptx 同一套设计）"""
    from app.services import deck_design
    return deck_design.render_slide_image(
        {"title": title, "bullets": bullets, "layout": "bullets"},
        out_abs, out_rel, name, page_no)


def _render_slide_only(slide: dict, out_abs: str, out_rel: str, name: str,
                       deck_title: str = "", audience: str = "") -> tuple[str, int]:
    """AI 生成的页面（无源文件）直接绘制版式化配图"""
    from app.services import deck_design
    return deck_design.render_slide_image(slide, out_abs, out_rel, name,
                                          int(slide.get("slide_index", 0)) + 1,
                                          deck_title=deck_title, audience=audience)


def _slide_to_image(slide, slide_index: int, out_abs: str, out_rel: str,
                    name: str) -> tuple[str, int]:
    """从已加载的 slide 对象生成页面图（不再重复解析整个 PPTX）"""
    title, bullets, _ = _extract_text_shapes(slide)
    return _draw_slide_image(title, bullets, out_abs, out_rel, name, slide_index + 1)


def convert_pptx_to_slides(pptx_path: str, courseware_id: int) -> list[dict[str, Any]]:
    """PPTX → 结构化页面列表（含文本层 + 瘦身图片）

    返回 [{slide_index, title, bullets, notes, tables, image_path, image_kb}]
    """
    from pptx import Presentation
    prs = Presentation(pptx_path)
    out_rel, out_abs = _slides_dir(courseware_id)

    results: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides):
        title, bullets, tables = _extract_text_shapes(slide)

        # 演讲者备注
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            pass

        # 优先用页面内最大的内嵌图作为该页视觉；没有则版式化渲染兜底
        image_path, image_kb = "", 0
        best_blob, best_size = None, 0
        for shape in slide.shapes:
            if shape.shape_type == 13 and getattr(shape, "image", None) is not None:  # PICTURE
                try:
                    blob = shape.image.blob
                    if len(blob) > best_size:
                        best_blob, best_size = blob, len(blob)
                except Exception:
                    continue
        if best_blob and best_size > MIN_IMAGE_BYTES:
            name = f"s{idx:03d}.jpg"
            image_path, image_kb = _compress_image(best_blob, out_abs, out_rel, name)

        if not image_path:
            name = f"s{idx:03d}.jpg"
            image_path, image_kb = _slide_to_image(slide, idx, out_abs, out_rel, name)

        results.append({
            "slide_index": idx, "title": title, "bullets": bullets,
            "notes": notes[:2000], "tables": tables,
            "image_path": image_path, "image_kb": image_kb,
        })
    return results


def layout_points(slide: dict[str, Any], include_core: bool = True) -> list[tuple[int, str]]:
    """把版式语义（卡片/对比/数据/图/流程/金句）摊平成「层级 + 文本」

    网页阅读器与 pptx 用版式渲染，而电子书与课堂讲义是线性 Markdown。
    这里做一次降维，避免这些内容在非版式产物里凭空消失。

    include_core=False 时不摊平 detail/example —— 讲义与 Word 会单独用
    「讲透」「案例」两个小标题排版，重复摊平会出现两遍。
    """
    out: list[tuple[int, str]] = []
    detail = str(slide.get("detail") or "").strip()
    if detail and include_core:
        out.append((0, detail))
    example = str(slide.get("example") or "").strip()
    if example and include_core:
        out.append((0, f"案例：{example}"))
    chart = slide.get("chart") or {}
    values = chart.get("values") or []
    if values:
        unit = str(chart.get("unit") or "")
        head = str(chart.get("series_name") or "数据").strip()
        out.append((0, head))
        for lb, v in zip(chart.get("labels") or [], values):
            out.append((1, f"{lb}：{v}{unit}"))
    for st in slide.get("steps") or []:
        title = str(st.get("title") or "").strip()
        desc = str(st.get("desc") or "").strip()
        if title:
            out.append((0, f"{title}：{desc}" if desc else title))
    for c in slide.get("cards") or []:
        title = str(c.get("title") or "").strip()
        desc = str(c.get("desc") or "").strip()
        if title:
            out.append((0, f"{title}：{desc}" if desc else title))
    comp = slide.get("compare") or {}
    for side in ("left", "right"):
        data = comp.get(side) or {}
        head = str(data.get("title") or "").strip()
        if head:
            out.append((0, head))
        for item in data.get("items") or []:
            text = str(item).strip()
            if text:
                out.append((1, text))
    for m in slide.get("metrics") or []:
        value = str(m.get("value") or "").strip()
        if not value:
            continue
        unit = str(m.get("unit") or "").strip()
        label = str(m.get("label") or "").strip()
        head = f"{value}{unit}"
        out.append((0, f"{head}｜{label}" if label else head))
    quote = str(slide.get("quote") or "").strip()
    if quote:
        out.append((0, f"「{quote}」"))
        by = str(slide.get("quote_by") or "").strip()
        if by:
            out.append((1, f"— {by}"))
    for t in slide.get("takeaway") or []:
        text = str(t).strip()
        if text:
            out.append((0, text))
    return out


def build_ebook_markdown(slides: list[dict[str, Any]], deck_title: str = "") -> str:
    """把散页拼成连贯读物（Markdown），这是「电子书体验」的核心"""
    lines = [f"# {deck_title}" if deck_title else "# 课件读物", ""]
    for s in slides:
        head = s.get("title") or f"第 {s['slide_index'] + 1} 页"
        lines.append(f"## {head}")
        lines.append("")
        bullets = s.get("bullets") or []
        if bullets:
            for b in bullets:
                lvl = min(int(b.get("level", 0) or 0), 3)
                lines.append(f"{'  ' * lvl}- {b.get('text', '')}")
            lines.append("")
        # 卡片/对比/数据/金句页的内容不在 bullets 里，摊平后补进来
        points = layout_points(s)
        for lvl, text in points:
            lines.append(f"{'  ' * min(lvl, 3)}- {text}")
        if points:
            lines.append("")
        for tbl in (s.get("tables") or [])[:1]:
            if tbl and len(tbl) > 1:
                lines.append("| " + " | ".join(tbl[0]) + " |")
                lines.append("|" + "|".join(["---"] * len(tbl[0])) + "|")
                for row in tbl[1:]:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")
        if s.get("notes"):
            lines.append(f"> 讲师备注：{s['notes']}")
            lines.append("")
    return "\n".join(lines)


def render_pptx(slides: list[dict[str, Any]], deck_title: str, out_path: str,
                audience: str = "") -> str:
    """结构化页面 → 16:9 原生 .pptx（版式由 deck_design 统一渲染）"""
    from app.services import deck_design
    return deck_design.render_pptx(slides, deck_title, out_path, audience=audience)


def estimate_optimization(original_bytes: int, slides: list[dict[str, Any]]) -> dict:
    """量化转换收益。关键指标是**首屏需下载量**而不是整包体积：
    阅读时只先取文本元数据（几 KB），图片按当前页懒加载，
    因此打开速度由 first_paint_kb 决定，而非整册体积。
    """
    text_kb = sum(len((s.get("title") or "") + "".join(b.get("text", "") for b in (s.get("bullets") or [])))
                  for s in slides) // 1024
    img_kb = sum(s.get("image_kb", 0) or 0 for s in slides)
    total_kb = img_kb + text_kb
    orig_kb = max(1, original_bytes // 1024)
    # 首屏 = 文本元数据 + 首页配图（其余页滑动到才拉）
    first_page_img = (slides[0].get("image_kb", 0) or 0) if slides else 0
    first_paint_kb = max(1, text_kb + first_page_img + len(slides))
    return {
        "original_kb": orig_kb,
        "ebook_kb": total_kb,
        "text_kb": text_kb,
        "image_kb": img_kb,
        "first_paint_kb": first_paint_kb,
        # 打开速度口径：首屏相比原整包下载省了多少
        "saved_percent": max(0, round((1 - first_paint_kb / orig_kb) * 100, 1)),
        # 全量阅读口径（把每页图都加载完的总量）
        "full_load_saved_percent": max(0, round((1 - total_kb / orig_kb) * 100, 1)),
        "lazy_loaded": True,
    }
