import pathlib

p = pathlib.Path(r'app/api/v1/multimodal.py')
content = p.read_text(encoding='utf-8')

# 1. 替换白名单
old_exts = '{".pdf", ".csv", ".txt", ".md", ".py", ".json"}'
new_exts = '{".pdf", ".csv", ".txt", ".md", ".py", ".json", ".xlsx", ".xls", ".docx"}'
content = content.replace(old_exts, new_exts)

# 2. 追加 _extract_text
if '_extract_text' not in content:
    func = '''

def _extract_text(save_path: str, ext: str) -> str:
    """根据文件类型提取文本内容"""
    if ext in (".txt", ".md", ".csv", ".json", ".py"):
        with open(save_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()[:50000]
    elif ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(save_path, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets[:3]:
            for row in ws.iter_rows(max_row=200, values_only=True):
                rows.append("\\t".join(str(c) if c is not None else "" for c in row))
        return "\\n".join(rows)[:50000]
    elif ext == ".docx":
        from docx import Document
        doc = Document(save_path)
        return "\\n".join(p.text for p in doc.paragraphs)[:50000]
    elif ext == ".pdf":
        import fitz
        text = []
        with fitz.open(save_path) as pdf:
            for page in pdf[:20]:
                text.append(page.get_text())
        return "\\n".join(text)[:50000]
    return "[无法提取文本内容]"
'''
    content += func

# 3. 插入调用
if 'extracted = _extract_text' not in content:
    content = content.replace(
        'f.write(content)',
        'f.write(content)\n    extracted = _extract_text(save_path, ext)'
    )

p.write_text(content, encoding='utf-8')
print('Done')
